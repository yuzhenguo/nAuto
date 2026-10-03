"""
review_manager.py
댓글목록.xlsx 읽기/쓰기 관리 (여러 기기 병행 작업 시 2중 작업 방지)

시트 구조 (헤더 자동 감지):
  Sheet1 (댓글):  검색어 | 댓글 | 작업여부
                  작업여부: 공백=미작업, W=작업중(예약), Y=완료, F=실패
  Sheet2 (계정):  폰ID | 순번 | 로그인아이디 | 작업완료여부
                  작업완료여부: 공백=미완료, Y=완료, F=실패(로그인/진입 실패)

모든 조회/기록은 매번 엑셀 파일을 새로 읽는다.
프로세스 내 락 + 파일 락(댓글목록.xlsx.lock)으로 "읽기 → W 기록 → 저장" 을
하나의 원자 구간으로 묶어 같은 댓글을 두 기기가 가져가지 않게 한다.
"""

import os
import re
import time
import threading
from contextlib import contextmanager
from typing import List, Optional, Tuple

import openpyxl


# ─── 헤더 키워드 (공백/대소문자 무시, 정확 일치 우선) ─────────────────────────
COMMENT_HEADERS = {
    "keyword": ["검색어", "상품명", "keyword"],
    "comment": ["댓글", "리뷰내용", "리뷰", "comment"],
    "status":  ["작업여부", "완료여부", "작업완료여부", "상태", "status"],
}
ACCOUNT_HEADERS = {
    "device_id": ["폰id", "기기id", "deviceid"],
    "seq":       ["순번", "번호"],
    "login_id":  ["로그인아이디", "로그인id", "아이디", "loginid"],
    "status":    ["작업완료여부", "완료여부", "작업여부", "상태", "status"],
}

STATUS_WORKING = "W"
STATUS_DONE = "Y"
STATUS_FAILED = "F"


class ExcelBusyError(Exception):
    """엑셀 파일 저장 불가 (엑셀 프로그램에서 열려 있는 경우 등)"""
    pass


def _norm(text) -> str:
    return re.sub(r"\s+", "", str(text or "")).lower()


def _cell_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def is_blank_status(val) -> bool:
    return _cell_str(val) == ""


def keyword_matches(keyword: str, texts: List[str]) -> bool:
    """
    댓글목록 검색어가 리뷰 카드의 상품명과 매칭되는지 판정.
    1) 공백 제거 후 검색어가 상품명에 포함
    2) 상품명이 검색어에 포함 (상품명 6자 이상일 때만)
    3) 검색어의 모든 단어가 카드 텍스트에 포함 (단어 순서가 다른 경우)
    """
    k = _norm(keyword)
    if not k:
        return False
    for t in texts:
        n = _norm(t)
        if not n:
            continue
        if k in n or (len(n) >= 6 and n in k):
            return True
    tokens = [_norm(x) for x in str(keyword).split() if _norm(x)]
    joined = _norm(" ".join(texts))
    return len(tokens) >= 2 and all(tok in joined for tok in tokens)


class CommentRow:
    def __init__(self, row_index: int, keyword: str, comment: str, status: str):
        self.row_index = row_index
        self.keyword = keyword
        self.comment = comment
        self.status = status

    def __repr__(self):
        return f"CommentRow(row={self.row_index}, keyword={self.keyword!r}, status={self.status!r})"


class AccountRow:
    def __init__(self, row_index: int, device_id: str, seq: str, login_id: str, status: str):
        self.row_index = row_index
        self.device_id = device_id
        self.seq = seq
        self.login_id = login_id
        self.status = status

    def __repr__(self):
        return f"AccountRow(row={self.row_index}, device={self.device_id!r}, login_id={self.login_id!r})"


class _FileLock:
    """다른 프로세스(프로그램 2중 실행)와의 동시 쓰기 방지용 잠금 파일"""

    def __init__(self, path: str, timeout: float = 60.0, stale_sec: float = 120.0):
        self.path = path
        self.timeout = timeout
        self.stale_sec = stale_sec

    def acquire(self):
        end = time.time() + self.timeout
        while True:
            try:
                fd = os.open(self.path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.write(fd, str(os.getpid()).encode())
                os.close(fd)
                return
            except FileExistsError:
                try:
                    if time.time() - os.path.getmtime(self.path) > self.stale_sec:
                        os.remove(self.path)  # 비정상 종료로 남은 잠금 파일 정리
                        continue
                except OSError:
                    pass
                if time.time() > end:
                    raise ExcelBusyError(f"엑셀 잠금 대기 시간 초과: {self.path}")
                time.sleep(0.1)

    def release(self):
        try:
            os.remove(self.path)
        except OSError:
            pass


def _detect_columns(ws, keywords: dict) -> dict:
    mapping = {}
    scores = {}
    for col_idx in range(1, ws.max_column + 1):
        cell = _norm(ws.cell(1, col_idx).value)
        if not cell:
            continue
        for field, kws in keywords.items():
            for kw in kws:
                kw_n = _norm(kw)
                if cell == kw_n:
                    score = 1000 + len(kw_n)
                elif kw_n in cell:
                    score = len(kw_n)
                else:
                    continue
                if score > scores.get(field, -1):
                    scores[field] = score
                    mapping[field] = col_idx
    return mapping


class ReviewManager:
    """댓글목록.xlsx 스레드/프로세스 안전 관리자 (매 호출마다 파일 새로 읽기)"""

    SAVE_RETRY_SEC = 30

    def __init__(self, xlsx_path: str):
        self.xlsx_path = xlsx_path
        self._lock = threading.RLock()
        self._file_lock = _FileLock(xlsx_path + ".lock")
        if not os.path.exists(xlsx_path):
            raise FileNotFoundError(xlsx_path)
        # 구조 검증 (시트/헤더 감지 실패 시 즉시 예외)
        with self._locked():
            self._open()

    # ─── 내부 ────────────────────────────────────────────────────────────────

    @contextmanager
    def _locked(self):
        with self._lock:
            self._file_lock.acquire()
            try:
                yield
            finally:
                self._file_lock.release()

    def _open(self):
        """워크북을 새로 읽고 (wb, 댓글시트, 댓글컬럼, 계정시트, 계정컬럼) 반환"""
        wb = openpyxl.load_workbook(self.xlsx_path)
        comment_ws = account_ws = None
        comment_cols = account_cols = None
        for ws in wb.worksheets:
            ac = _detect_columns(ws, ACCOUNT_HEADERS)
            if account_ws is None and "login_id" in ac and "device_id" in ac:
                account_ws, account_cols = ws, ac
                continue
            cc = _detect_columns(ws, COMMENT_HEADERS)
            if comment_ws is None and "keyword" in cc and "comment" in cc:
                comment_ws, comment_cols = ws, cc

        if comment_ws is None:
            raise ValueError("댓글 시트(검색어/댓글 헤더)를 찾을 수 없습니다.")
        if account_ws is None:
            raise ValueError("계정 시트(폰ID/로그인아이디 헤더)를 찾을 수 없습니다.")

        # 상태 컬럼이 없으면 맨 끝에 생성
        if "status" not in comment_cols:
            col = comment_ws.max_column + 1
            comment_ws.cell(1, col).value = "작업여부"
            comment_cols["status"] = col
        if "status" not in account_cols:
            col = account_ws.max_column + 1
            account_ws.cell(1, col).value = "작업완료여부"
            account_cols["status"] = col
        return wb, comment_ws, comment_cols, account_ws, account_cols

    def _save(self, wb):
        end = time.time() + self.SAVE_RETRY_SEC
        last_err = None
        while time.time() < end:
            try:
                wb.save(self.xlsx_path)
                return
            except PermissionError as e:
                last_err = e
                time.sleep(1.0)
        raise ExcelBusyError(
            f"엑셀 저장 실패 (엑셀 프로그램에서 파일이 열려 있으면 닫아주세요): {last_err}"
        )

    def _iter_comment_rows(self, ws, cols):
        for r in range(2, ws.max_row + 1):
            keyword = _cell_str(ws.cell(r, cols["keyword"]).value)
            comment = _cell_str(ws.cell(r, cols["comment"]).value)
            if not keyword or not comment:
                continue
            yield CommentRow(r, keyword, comment, _cell_str(ws.cell(r, cols["status"]).value))

    def _iter_account_rows(self, ws, cols):
        for r in range(2, ws.max_row + 1):
            login_id = _cell_str(ws.cell(r, cols["login_id"]).value)
            if not login_id:
                continue
            yield AccountRow(
                r,
                _cell_str(ws.cell(r, cols["device_id"]).value),
                _cell_str(ws.cell(r, cols["seq"]).value) if "seq" in cols else "",
                login_id,
                _cell_str(ws.cell(r, cols["status"]).value),
            )

    # ─── 댓글 (Sheet1) ───────────────────────────────────────────────────────

    def claim_comment(self, candidates: List[List[str]]) -> Tuple[Optional[int], Optional[CommentRow]]:
        """
        리뷰 버튼 후보(위에서부터 순서대로, 각 후보는 카드 텍스트 목록) 중
        작업여부 공백인 댓글과 처음으로 매칭되는 후보를 찾아 해당 댓글을 W로 예약.
        Returns: (후보 인덱스, CommentRow) / 매칭 없음: (None, None)
        """
        with self._locked():
            wb, cws, ccols, _, _ = self._open()
            pending = [r for r in self._iter_comment_rows(cws, ccols) if is_blank_status(r.status)]
            if not pending:
                return None, None
            for idx, texts in enumerate(candidates):
                for row in pending:
                    if keyword_matches(row.keyword, texts):
                        cws.cell(row.row_index, ccols["status"]).value = STATUS_WORKING
                        self._save(wb)
                        row.status = STATUS_WORKING
                        return idx, row
            return None, None

    def set_comment_status(self, row_index: int, status: str):
        with self._locked():
            wb, cws, ccols, _, _ = self._open()
            cws.cell(row_index, ccols["status"]).value = status or None
            self._save(wb)

    def release_comment(self, row_index: int) -> bool:
        """작업중(W)인 댓글을 공백으로 되돌림 (리뷰 미작성으로 끝난 경우). Y/F는 건드리지 않음."""
        with self._locked():
            wb, cws, ccols, _, _ = self._open()
            cell = cws.cell(row_index, ccols["status"])
            if _cell_str(cell.value).upper() != STATUS_WORKING:
                return False
            cell.value = None
            self._save(wb)
            return True

    # ─── 계정 (Sheet2) ───────────────────────────────────────────────────────

    def next_pending_account(self, device_id: str) -> Optional[AccountRow]:
        """해당 폰ID 의 작업완료여부 공백인 첫 번째 로그인 아이디"""
        dev = _norm(device_id)
        with self._locked():
            _, _, _, aws, acols = self._open()
            for row in self._iter_account_rows(aws, acols):
                if _norm(row.device_id) == dev and is_blank_status(row.status):
                    return row
        return None

    def set_account_status(self, row_index: int, status: str):
        with self._locked():
            wb, _, _, aws, acols = self._open()
            aws.cell(row_index, acols["status"]).value = status or None
            self._save(wb)

    def get_device_summary(self) -> dict:
        """{폰ID(대문자): {"seq": 순번, "has_pending": 미완료 아이디 존재 여부}}"""
        result = {}
        with self._locked():
            _, _, _, aws, acols = self._open()
            for row in self._iter_account_rows(aws, acols):
                key = row.device_id.strip().upper()
                if not key:
                    continue
                info = result.setdefault(key, {"seq": row.seq, "has_pending": False})
                if not info["seq"] and row.seq:
                    info["seq"] = row.seq
                if is_blank_status(row.status):
                    info["has_pending"] = True
        return result
